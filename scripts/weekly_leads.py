#!/usr/bin/env python3
"""
Weekly 3-step lead pipeline for the Nemko CRA lead tracker.

Each week produces a new workbook in "Weekly Leads/" with three sheets:
  1. ALL Weekly Leads — every new US lead scraped from the FCC EAS search
  2. CRA Scoring     — gates/questions/grade per lead (research fills this)
  3. Contacts        — contact info (from FCC Form 731) for A/B-grade leads

Subcommands (paths relative to the "Nemko Lead Gen" folder):
  ingest scraped.json          create weekly workbook + append new leads to
                               the master ledger; prints JSON of the new
                               leads' Form-731 URLs for the browser to fetch
  enrich731 --week F forms.json  fill product + contact columns from parsed
                               Form 731 data (weekly sheet 1 + master Product)
  finalize --week F            copy A/B-grade rows' contacts to the Contacts
                               sheet and sync scores back to the master

The master workbook ("Nemko CRA Leads.xlsx") stays the cumulative dedupe
ledger. "Nemko_CRA_DATA_Transfer.xlsx" (internal leads) is read-only, always.
No email is ever sent by this script.
"""
import argparse
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

sys.path.insert(0, str(Path(__file__).resolve().parent))
import fcc_pull
from fcc_pull import (COL, CLASS_INFO, GRADE_BANDS, Q_POINTS, SCRIPT_TRUE_QS,
                      class_code, grade_for, internal_companies, linkedin_url,
                      norm_company, read_config)

FOLDER = Path(__file__).resolve().parent.parent
MASTER = FOLDER / "Nemko CRA Leads.xlsx"

ARIAL = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=ARIAL, bold=True, color="FFFFFF", size=10)
LINK_FONT = Font(name=ARIAL, size=10, color="0563C1", underline="single")
MAX_ROW = 1000

ALL_SHEET = "ALL Weekly Leads"
SCORE_SHEET = "CRA Scoring"
CONTACT_SHEET = "Contacts"

ALL_COLS = ["Lead ID", "Company", "City", "State", "FCC ID", "Grant Date",
            "Equipment Class", "Connectivity", "Product", "Application Purpose",
            "Contact Name", "Contact Title", "Contact Phone", "Contact Email",
            "Form 731", "Exhibits"]
ALL_W = [9, 30, 15, 8, 17, 11, 10, 14, 34, 16, 20, 22, 15, 26, 9, 9]

# Scoring sheet reuses the master's G/Q header text (cols 7..26 of All Leads)
CONTACT_COLS = ["Company", "Product", "GRADE", "TOTAL /100", "Contact Name",
                "Contact Title", "Contact Phone", "Contact Email", "FCC ID", "Notes"]
CONTACT_W = [30, 34, 16, 10, 20, 24, 15, 28, 17, 30]

GRADE_LIST = "A — Pursue now,B — Strong,C — Qualify,D — Deprioritize,DQ (gate)"


def style_header(ws, names, widths):
    for i, name in enumerate(names, 1):
        c = ws.cell(row=1, column=i, value=name)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1] if i <= len(widths) else 10
    ws.row_dimensions[1].height = 40
    ws.freeze_panes = "C2"


def col_index(names, name):
    return names.index(name) + 1


def qualify_prefixes(cfg):
    grades = [g.strip().upper() for g in str(cfg.get("qualify_grades") or "A,B").split(",") if g.strip()]
    return tuple(grades)


def contact_name(rec):
    return " ".join(x for x in (rec.get("contact_first", ""), rec.get("contact_last", "")) if x).strip()


# ------------------------------------------------------------------ ingest
def cmd_ingest(args):
    data = json.loads(Path(args.scraped).read_text())
    rows, window = data["rows"], data.get("window", {})
    w_from = window.get("from", date.today().isoformat())
    w_to = window.get("to", date.today().isoformat())
    today = date.today()

    master = load_workbook(MASTER)
    cfg = read_config(master)
    ws = master["All Leads"]
    master_headers = [c.value for c in ws[1]]
    gq_headers = master_headers[6:26]  # G1..G6, DISQUALIFIED?, Q1..Q13

    # dedupe state (same logic as fcc_pull.main)
    seen_ids, recent = set(), {}
    dedupe_days = int(cfg.get("company_dedupe_days") or 30)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) >= COL["fcc_id"] and row[COL["fcc_id"] - 1]:
            seen_ids.add(str(row[COL["fcc_id"] - 1]).strip().upper())
        comp = norm_company(row[COL["company"] - 1] if len(row) >= COL["company"] else "")
        added = row[COL["date_added"] - 1] if len(row) >= COL["date_added"] else None
        if comp and added:
            try:
                d = added.date() if isinstance(added, datetime) else datetime.strptime(str(added)[:10], "%Y-%m-%d").date()
            except ValueError:
                continue
            recent[comp] = max(recent.get(comp, d), d)
    internal = internal_companies(cfg)
    blacklist = fcc_pull.blacklisted_companies(master)

    max_id = 0
    for row in ws.iter_rows(min_row=2, min_col=COL["lead_id"], max_col=COL["lead_id"], values_only=True):
        m = re.match(r"L(\d+)", str(row[0] or ""))
        if m:
            max_id = max(max_id, int(m.group(1)))

    us = fcc_pull.is_us
    base_total = sum(p for i, p in enumerate(Q_POINTS, 1) if i in SCRIPT_TRUE_QS)
    new_leads, skipped_dup, skipped_internal, skipped_nonus, skipped_blacklist = [], 0, 0, 0, 0
    for r in rows:
        if not us(r.get("country", "United States")):
            skipped_nonus += 1
            continue
        fid = str(r["fcc_id"]).strip().upper()
        comp = norm_company(r["company"])
        company_dup = (not getattr(args, "include_repeat_filers", False)
                       and comp in recent and (today - recent[comp]).days <= dedupe_days)
        if fid in seen_ids or company_dup:
            skipped_dup += 1
            continue
        if comp in internal:
            skipped_internal += 1
            continue
        if comp in blacklist:
            skipped_blacklist += 1
            continue
        seen_ids.add(fid)
        recent[comp] = today
        max_id += 1
        r = dict(r)
        r["lead_id"] = f"L{max_id:04d}"
        new_leads.append(r)

    # ---- weekly workbook
    weekly_dir = FOLDER / str(cfg.get("weekly_folder") or "Weekly Leads")
    weekly_dir.mkdir(exist_ok=True)
    label = args.label or f"Week of {w_from} to {w_to}"
    week_path = weekly_dir / f"{label}.xlsx"
    if week_path.exists() and not args.force:
        sys.exit(f"ERROR: {week_path.name} already exists (use --force to overwrite)")

    wk = Workbook()
    wa = wk.active
    wa.title = ALL_SHEET
    style_header(wa, ALL_COLS, ALL_W)
    body = Font(name=ARIAL, size=10)
    for i, r in enumerate(new_leads, 2):
        code = class_code(r.get("equipment_class", ""))
        vals = [r["lead_id"], r["company"], r.get("city", "").rstrip(","), r.get("state", ""),
                r["fcc_id"], r.get("grant_date", ""), code,
                CLASS_INFO.get(code, "Unknown"), "", r.get("purpose", ""),
                "", "", "", "", "View Form", "Detail"]
        for c, v in enumerate(vals, 1):
            cell = wa.cell(row=i, column=c, value=v)
            cell.font = body
        for cname, url_key in (("Form 731", "form731"), ("Exhibits", "exhibits")):
            url = r.get(url_key)
            cell = wa.cell(row=i, column=col_index(ALL_COLS, cname))
            if url:
                cell.hyperlink = url
                cell.font = LINK_FONT
            else:
                cell.value = ""

    sc = wk.create_sheet(SCORE_SHEET)
    score_cols = ["Company", "FCC ID"] + gq_headers + ["TOTAL /100", "GRADE",
                  "Key evidence", "Confidence", "Research notes"]
    style_header(sc, score_cols, [30, 17] + [7] * 20 + [10, 16, 40, 11, 40])
    for i, r in enumerate(new_leads, 2):
        sc.cell(row=i, column=1, value=r["company"]).font = body
        sc.cell(row=i, column=2, value=r["fcc_id"]).font = body
        for gq in range(6):                      # G1..G6
            sc.cell(row=i, column=3 + gq, value=False)
        sc.cell(row=i, column=9, value="OK")     # DISQUALIFIED?
        for q in range(13):                      # Q1..Q13
            sc.cell(row=i, column=10 + q, value=(q + 1) in SCRIPT_TRUE_QS)
        sc.cell(row=i, column=23, value=base_total)
        sc.cell(row=i, column=24, value=grade_for(base_total))
        sc.cell(row=i, column=26, value="Low")
    n = len(score_cols)
    dv_bool = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
    dv_bool.add(f"C2:H{MAX_ROW}")
    dv_bool.add(f"J2:V{MAX_ROW}")
    dv_dq = DataValidation(type="list", formula1='"OK,DQ"', allow_blank=True)
    dv_dq.add(f"I2:I{MAX_ROW}")
    dv_grade = DataValidation(type="list", formula1=f'"{GRADE_LIST}"', allow_blank=True)
    dv_grade.add(f"X2:X{MAX_ROW}")
    dv_conf = DataValidation(type="list", formula1='"High,Med,Low"', allow_blank=True)
    dv_conf.add(f"Z2:Z{MAX_ROW}")
    for dv in (dv_bool, dv_dq, dv_grade, dv_conf):
        sc.add_data_validation(dv)
    for prefix, color in (('"A ', "C6EFCE"), ('"B ', "E2EFDA"), ('"C ', "FFEB9C"),
                          ('"D ', "FCE4D6"), ('"DQ', "FFC7CE")):
        sc.conditional_formatting.add(f"X2:X{MAX_ROW}", FormulaRule(
            formula=[f'LEFT(X2,{len(prefix) - 1})={prefix}"'],
            fill=PatternFill("solid", fgColor=color)))

    ct = wk.create_sheet(CONTACT_SHEET)
    style_header(ct, CONTACT_COLS, CONTACT_W)
    ct.cell(row=2, column=1,
            value="(filled by finalize once CRA Scoring grades are set — A/B leads only)"
            ).font = Font(name=ARIAL, size=10, italic=True, color="808080")

    wk.save(week_path)

    # ---- append to master ledger (same shape as fcc_pull appends)
    for r in new_leads:
        code = class_code(r.get("equipment_class", ""))
        gd = fcc_pull.parse_grant_date(r.get("grant_date", ""))
        row_i = ws.max_row + 1
        office = ", ".join(x for x in (r.get("city", "").rstrip(","), r.get("state", "")) if x)
        evidence = (f"{CLASS_INFO.get(code, 'radio device')}; "
                    f"FCC grant {gd.isoformat() if gd else r.get('grant_date', '')} ({r['fcc_id']})")
        values = {
            "status": "Not Contacted", "company": r["company"], "office": office,
            "total": base_total, "grade": grade_for(base_total),
            "key_evidence": evidence, "confidence": "Low",
            "lead_id": r["lead_id"], "date_added": today.isoformat(),
            "source": "FCC", "fcc_id": r["fcc_id"],
            "grant_date": gd.isoformat() if gd else r.get("grant_date", ""),
            "equipment_class": code or r.get("equipment_class", ""),
            "product": "", "country": r.get("country", "United States"),
        }
        for key, v in values.items():
            ws.cell(row=row_i, column=COL[key], value=v)
        for c in range(fcc_pull.G_FIRST, fcc_pull.G_LAST + 1):
            ws.cell(row=row_i, column=c, value=False)
        ws.cell(row=row_i, column=fcc_pull.DQ_COL, value="OK")
        for qi, c in enumerate(range(fcc_pull.Q_FIRST, fcc_pull.Q_LAST + 1), 1):
            ws.cell(row=row_i, column=c, value=qi in SCRIPT_TRUE_QS)
        link = ws.cell(row=row_i, column=COL["linkedin"], value="Find PM")
        link.hyperlink = linkedin_url(r["company"])
        link.font = LINK_FONT
    if new_leads:
        ws.tables["LeadsTable"].ref = f"A1:AL{ws.max_row}"
    for row in master["Config"].iter_rows(min_row=2, max_col=2):
        if str(row[0].value).strip() == "last_run_date":
            row[1].value = w_to
            break
    try:
        master.save(MASTER)
    except PermissionError:
        sys.exit("ERROR: could not save master — close 'Nemko CRA Leads.xlsx' in Excel and rerun.")

    print(json.dumps({
        "week_file": str(week_path),
        "new": len(new_leads), "dup": skipped_dup,
        "internal": skipped_internal, "non_us": skipped_nonus,
        "blacklist": skipped_blacklist,
        "fetch_731": [{"fcc_id": r["fcc_id"], "form731": r.get("form731", "")}
                      for r in new_leads],
    }, indent=1))


# --------------------------------------------------------------- enrich731
def cmd_enrich731(args):
    forms = json.loads(Path(args.forms).read_text())
    by_id = {str(f["fcc_id"]).strip().upper(): f for f in forms}
    wk = load_workbook(args.week)
    wa = wk[ALL_SHEET]
    c_fcc = col_index(ALL_COLS, "FCC ID")
    filled = 0
    for row in wa.iter_rows(min_row=2):
        fid = str(row[c_fcc - 1].value or "").strip().upper()
        rec = by_id.get(fid)
        if not rec:
            continue
        sets = {
            "Product": rec.get("product", ""),
            "Contact Name": contact_name(rec),
            "Contact Title": rec.get("contact_title", ""),
            "Contact Phone": rec.get("contact_phone", ""),
            "Contact Email": rec.get("contact_email", ""),
        }
        for cname, v in sets.items():
            if v:
                row[col_index(ALL_COLS, cname) - 1].value = v
        filled += 1
    wk.save(args.week)

    # product also goes to the master (blank cells only)
    master = load_workbook(MASTER)
    ws = master["All Leads"]
    m_filled = 0
    for row in ws.iter_rows(min_row=2):
        fid = str(row[COL["fcc_id"] - 1].value or "").strip().upper()
        rec = by_id.get(fid)
        if rec and rec.get("product") and not row[COL["product"] - 1].value:
            row[COL["product"] - 1].value = rec["product"]
            m_filled += 1
    try:
        master.save(MASTER)
    except PermissionError:
        sys.exit("ERROR: could not save master — close 'Nemko CRA Leads.xlsx' in Excel and rerun.")
    print(f"731 data: {filled} weekly rows enriched, {m_filled} master products filled "
          f"({len(forms)} form records supplied)")


# ---------------------------------------------------------------- finalize
def cmd_finalize(args):
    wk = load_workbook(args.week)
    wa, sc, ct = wk[ALL_SHEET], wk[SCORE_SHEET], wk[CONTACT_SHEET]
    master = load_workbook(MASTER)
    ws = master["All Leads"]
    cfg = read_config(master)
    prefixes = qualify_prefixes(cfg)

    all_by_id = {}
    for row in wa.iter_rows(min_row=2):
        fid = str(row[col_index(ALL_COLS, "FCC ID") - 1].value or "").strip().upper()
        if fid:
            all_by_id[fid] = row
    master_by_id = {}
    for row in ws.iter_rows(min_row=2):
        fid = str(row[COL["fcc_id"] - 1].value or "").strip().upper()
        if fid:
            master_by_id[fid] = row

    # clear Contacts data rows (idempotent)
    if ct.max_row > 1:
        ct.delete_rows(2, ct.max_row - 1)

    body = Font(name=ARIAL, size=10)
    n_contacts, n_synced, n_dq = 0, 0, 0
    out_row = 2
    scored = sorted(
        (r for r in sc.iter_rows(min_row=2) if r[1].value),
        key=lambda r: -(r[22].value or 0))
    for r in scored:
        fid = str(r[1].value or "").strip().upper()
        grade = str(r[23].value or "")
        total = r[22].value
        if grade.startswith("DQ"):
            n_dq += 1
        # sync to master
        m = master_by_id.get(fid)
        if m is not None:
            for j in range(20):  # G1..Q13 block (master cols 7..26 == score cols 3..22)
                m[6 + j].value = r[2 + j].value
            m[COL["total"] - 1].value = total
            m[COL["grade"] - 1].value = grade
            if r[24].value:
                m[COL["key_evidence"] - 1].value = r[24].value
            if r[25].value:
                m[COL["confidence"] - 1].value = r[25].value
            n_synced += 1
        # contacts for qualifying grades
        if any(grade.startswith(p) for p in prefixes):
            a = all_by_id.get(fid)
            if a is None:
                continue
            def av(name):
                return a[col_index(ALL_COLS, name) - 1].value or ""
            vals = [av("Company"), av("Product"), grade, total, av("Contact Name"),
                    av("Contact Title"), av("Contact Phone"), av("Contact Email"),
                    fid, ""]
            for c, v in enumerate(vals, 1):
                ct.cell(row=out_row, column=c, value=v).font = body
            out_row += 1
            n_contacts += 1

    wk.save(args.week)
    try:
        master.save(MASTER)
    except PermissionError:
        sys.exit("ERROR: could not save master — close 'Nemko CRA Leads.xlsx' in Excel and rerun.")
    print(f"finalize: {n_contacts} contacts on '{CONTACT_SHEET}' (grades {'/'.join(prefixes)}), "
          f"{n_dq} DQ'd, {n_synced} rows synced to master")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("ingest", help="scraped.json -> weekly workbook + master append")
    p.add_argument("scraped")
    p.add_argument("--label", help="weekly file label (default: Week of <from> to <to>)")
    p.add_argument("--force", action="store_true", help="overwrite an existing weekly file")
    p.add_argument("--include-repeat-filers", action="store_true",
                   help="keep filings from companies already seen recently (disables the "
                        "company-within-N-days dedupe; FCC-ID, internal-list and blacklist "
                        "skips still apply)")
    p.set_defaults(func=cmd_ingest)
    p = sub.add_parser("enrich731", help="forms.json -> fill product/contact columns")
    p.add_argument("--week", required=True)
    p.add_argument("forms")
    p.set_defaults(func=cmd_enrich731)
    p = sub.add_parser("finalize", help="grades -> Contacts sheet + master sync")
    p.add_argument("--week", required=True)
    p.set_defaults(func=cmd_finalize)
    args = ap.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
